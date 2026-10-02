"""The HTTP service, vision off."""
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from drawcheck import api, samples

pytestmark = pytest.mark.skipif(samples.symbol_font() is None, reason="no font with GD&T glyphs")


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "DATA", tmp_path / "runs")
    monkeypatch.setattr(api, "TOKEN", "")
    return TestClient(api.app)


@pytest.fixture(scope="module")
def pdf_bytes(tmp_path_factory):
    p = samples.make("defective", tmp_path_factory.mktemp("s") / "defective.pdf")
    return p.read_bytes()


def _upload(client, data, customer="acme"):
    r = client.post("/api/runs", files={"file": ("defective.pdf", data, "application/pdf")},
                    data={"customer": customer, "vision": "off"})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_full_review_cycle(client, pdf_bytes):
    rid = _upload(client, pdf_bytes)
    run = client.get(f"/api/runs/{rid}").json()        # background task already ran
    assert run["status"] == "done", run
    assert run["errors"] > 0 and run["drawing_number"] == "OF-1002"
    findings = run["report"]["findings"]

    png = client.get(f"/api/runs/{rid}/pages/0.png")
    assert png.status_code == 200 and png.content[:4] == b"\x89PNG"

    th1 = next(f for f in findings if f["rule"] == "TH-001")
    r = client.post(f"/api/runs/{rid}/findings/{th1['id']}", json={"decision": "accepted"})
    assert r.status_code == 200
    after = client.get(f"/api/runs/{rid}").json()["report"]
    assert next(f for f in after["findings"] if f["id"] == th1["id"])["decision"] == "accepted"

    xlsx = client.get(f"/api/runs/{rid}/files/queries.xlsx")
    rows = list(load_workbook(io.BytesIO(xlsx.content))["Technical Queries"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(findings) - 1                  # the accepted deviation is not a query any more
    assert all(row[5] != "TH-001" for row in rows)

    # Same drawing resubmitted: the decision is shown, and it is still not a query.
    rid2 = _upload(client, pdf_bytes)
    rep2 = client.get(f"/api/runs/{rid2}").json()["report"]
    assert next(f for f in rep2["findings"] if f["rule"] == "TH-001")["decision"] == "accepted"
    assert client.get("/api/reviews/stats").json()["TH-001"] == {"accepted": 1}

    assert client.delete(f"/api/runs/{rid}").status_code == 200
    assert client.get(f"/api/runs/{rid}").status_code == 404


def test_rejects_non_pdf(client):
    r = client.post("/api/runs", files={"file": ("x.pdf", b"hello", "application/pdf")}, data={"vision": "off"})
    assert r.status_code == 415


def test_token(client, pdf_bytes, monkeypatch):
    monkeypatch.setattr(api, "TOKEN", "s3cret")
    assert client.get("/api/runs").status_code == 401
    assert client.get("/api/runs", headers={"Authorization": "Bearer s3cret"}).status_code == 200
    assert client.get("/").status_code == 200              # the page itself is static


def test_bad_ids(client):
    assert client.get("/api/runs/../../etc").status_code in (400, 404)
    assert client.get("/api/runs/nope123").status_code == 404
