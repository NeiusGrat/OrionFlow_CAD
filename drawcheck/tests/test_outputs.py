"""Report files and review memory."""
import json

import pymupdf
import pytest
from openpyxl import load_workbook

from drawcheck import check_file, samples
from drawcheck.report import write_annotated_pdf, write_json, write_query_xlsx
from drawcheck.store import Store

pytestmark = pytest.mark.skipif(samples.symbol_font() is None, reason="no font with GD&T glyphs")


@pytest.fixture(scope="module")
def defective(tmp_path_factory):
    d = tmp_path_factory.mktemp("out")
    return samples.make("defective", d / "defective.pdf")


def test_outputs(defective, tmp_path):
    report, drawing = check_file(defective, vision="off")
    write_json(report, tmp_path / "r.json")
    assert json.loads((tmp_path / "r.json").read_text(encoding="utf-8"))["errors"] == report.count("error")

    write_annotated_pdf(report, defective, tmp_path / "a.pdf")
    with pymupdf.open(str(tmp_path / "a.pdf")) as doc:
        assert len(doc) >= 2                                   # drawing + summary
        boxed = [f for f in report.findings if f.bbox]
        assert len([a for a in doc[0].annots() if a.type[1] == "Square"]) == len(boxed)
        assert "Drawing check" in doc[1].get_text()

    sizes = {p.index: (p.width, p.height) for p in drawing.pages}
    write_query_xlsx(report, tmp_path / "q.xlsx", sizes)
    ws = load_workbook(tmp_path / "q.xlsx")["Technical Queries"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(report.findings)
    assert rows[0][0] == "TQ-001" and rows[0][1] == "OF-1002"


def test_fingerprints_are_stable(defective):
    a, _ = check_file(defective, vision="off")
    b, _ = check_file(defective, vision="off")
    assert [f.id for f in a.findings] == [f.id for f in b.findings]
    assert len({f.id for f in a.findings}) == len(a.findings)


def test_store_accept_and_reject(defective, tmp_path):
    store = Store(tmp_path / "s.sqlite")
    report, _ = check_file(defective, vision="off", store=store, customer="acme")
    fp = next(f for f in report.findings if f.rule == "TH-001")     # accepted deviation
    misread = next(f for f in report.findings if f.rule == "SF-002")  # pretend the checker was wrong
    store.decide(fp, "accepted", customer="acme")
    store.decide(misread, "rejected", customer="acme")

    again, _ = check_file(defective, vision="off", store=store, customer="acme")
    open_rules = {f.rule: f for f in again.findings}
    assert open_rules["TH-001"].decision == "accepted"            # same drawing: shown as accepted
    assert "SF-002" not in open_rules                              # rejected: moved out
    assert any(f.rule == "SF-002" for f in again.suppressed)

    other, _ = check_file(defective, vision="off", store=store, customer="other-customer")
    assert {f.rule for f in other.findings} >= {"TH-001", "SF-002"}   # decisions are per customer
    assert store.stats()["TH-001"] == {"accepted": 1}


def test_accepted_deviation_carries_to_other_drawings(tmp_path):
    from drawcheck.model import Finding, Report

    store = Store(tmp_path / "s.sqlite")
    first = Finding("TH-001", "warning", "t", "m", "q", page=0, bbox=(10, 10, 50, 20),
                    evidence="1/4-20 UNC", id="aaa")
    store.decide(first, "accepted", customer="acme")
    # A different drawing: other position, other fingerprint, same callout text.
    other = Finding("TH-001", "warning", "t", "m", "q", page=1, bbox=(300, 300, 340, 310),
                    evidence="1/4-20  unc", id="bbb")
    unrelated = Finding("TH-001", "warning", "t", "m", "q", page=0, evidence="3/8-16 UNC", id="ccc")
    r = Report("x.pdf", 2, findings=[other, unrelated])
    store.apply(r, "acme")
    assert [f.id for f in r.suppressed] == ["bbb"] and r.suppressed[0].decision == "accepted"
    assert [f.id for f in r.findings] == ["ccc"]
    r2 = Report("x.pdf", 2, findings=[Finding("TH-001", "warning", "t", "m", "q", evidence="1/4-20 UNC", id="bbb")])
    store.apply(r2, "someone-else")
    assert len(r2.findings) == 1                      # deviations are per customer
