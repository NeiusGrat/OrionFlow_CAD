"""Report outputs: JSON, an annotated copy of the drawing, and the technical
query list a supplier sends back to the customer (Excel)."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pymupdf

from .model import Finding, Report

SEVERITY_RGB = {"error": (0.85, 0.1, 0.1), "warning": (0.95, 0.55, 0.0), "info": (0.15, 0.4, 0.85)}


def fingerprint(f: Finding) -> str:
    """Stable id: the same finding on a re-submitted drawing gets the same id."""
    box = tuple(round(v / 5) for v in f.bbox) if f.bbox else ()
    key = f"{f.rule}|{f.page}|{f.evidence}|{box}"
    return hashlib.sha1(key.encode("utf-8")).hexdigest()[:10]


def zone(f: Finding, width: float, height: float, cols: int = 8, rows: int = 6) -> str:
    """Approximate drawing zone (ISO 5457 style: rows lettered top-down, columns numbered)."""
    if not f.bbox:
        return ""
    cx, cy = (f.bbox[0] + f.bbox[2]) / 2, (f.bbox[1] + f.bbox[3]) / 2
    c = min(int(cx / width * cols), cols - 1) + 1
    r = "ABCDEFGH"[min(int(cy / height * rows), rows - 1)]
    return f"{r}{c}"


def write_json(report: Report, path: str | Path) -> Path:
    path = Path(path)
    path.write_text(json.dumps(report.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def write_annotated_pdf(report: Report, src: str | Path, out: str | Path) -> Path:
    out = Path(out)
    with pymupdf.open(str(src)) as doc:
        for n, f in enumerate(report.findings, 1):
            if f.page is None or f.bbox is None or f.page >= len(doc):
                continue
            page = doc[f.page]
            rgb = SEVERITY_RGB[f.severity]
            # Finding boxes are in the page as displayed; annotations take unrotated coordinates.
            back = page.derotation_matrix
            r = pymupdf.Rect(f.bbox) + (-3, -3, 3, 3)
            a = page.add_rect_annot(r * back)
            a.set_colors(stroke=rgb)
            a.set_border(width=1.5)
            a.set_info(title=f"{f.rule} {f.severity}", content=f"{f.title}\n\n{f.query}")
            a.update()
            tag = pymupdf.Rect(r.x1 + 1, r.y0 - 11, r.x1 + 25, r.y0 + 1) * back
            t = page.add_freetext_annot(tag, str(n), fontsize=8, text_color=(1, 1, 1), fill_color=rgb,
                                        rotate=page.rotation, align=1)
            t.update()
        _summary_pages(doc, report)
        doc.save(str(out), garbage=3, deflate=True)
    return out


def _summary_pages(doc: pymupdf.Document, report: Report) -> None:
    lines = [
        f"Drawing check: {Path(report.source).name}",
        f"Drawing no: {report.title.get('drawing_number') or report.title.get('part_number') or '-'}   "
        f"Rev: {report.title.get('revision') or '-'}   Material: {report.title.get('material') or '-'}",
        f"{report.count('error')} errors   {report.count('warning')} warnings   {report.count('info')} info",
        "",
    ]
    for w in report.reader_warnings:
        lines.append(f"Reader: {w}")
    if report.reader_warnings:
        lines.append("")
    for n, f in enumerate(report.findings, 1):
        where = f"p{f.page + 1}" if f.page is not None else "drawing"
        conf = "" if f.confidence in ("vector", "confirmed", "absence") else f"  [{f.confidence}: verify]"
        lines.append(f"{n:>3}. [{f.severity.upper()}] {f.rule} {f.title} ({where}){conf}")
        lines.append(f"      Q: {f.query}")
        if f.standard:
            lines.append(f"      Ref: {f.standard}" + (f" cl. {f.clause}" if f.clause else ""))
    from .samples import symbol_font

    font_file = symbol_font()            # GD&T glyphs; Helvetica prints them as '?'
    font = "sym" if font_file else "helv"
    w, h = pymupdf.paper_size("a4")
    rect = pymupdf.Rect(36, 36, w - 36, h - 36)

    def new_page(d: pymupdf.Document) -> pymupdf.Page:
        p = d.new_page(width=w, height=h)
        if font_file:
            p.insert_font(fontname="sym", fontfile=font_file)
        return p

    def fits(n: int) -> bool:
        probe = pymupdf.open()
        try:
            return new_page(probe).insert_textbox(rect, "\n".join(remaining[:n]), fontsize=8, fontname=font) >= 0
        finally:
            probe.close()

    remaining = lines
    while remaining:
        lo, hi, fit = 1, len(remaining), 1     # as many lines as fit on one page
        while lo <= hi:
            mid = (lo + hi) // 2
            if fits(mid):
                fit, lo = mid, mid + 1
            else:
                hi = mid - 1
        new_page(doc).insert_textbox(rect, "\n".join(remaining[:fit]), fontsize=8, fontname=font)
        remaining = remaining[fit:]


def write_query_xlsx(report: Report, out: str | Path, page_sizes: dict[int, tuple[float, float]] | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    out = Path(out)
    wb = Workbook()
    ws = wb.active
    ws.title = "Technical Queries"
    head = ["TQ No", "Drawing No", "Rev", "Sheet", "Zone", "Rule", "Severity", "Query to customer",
            "Reference", "Customer response", "Status"]
    ws.append(head)
    dn = report.title.get("drawing_number") or report.title.get("part_number") or ""
    rev = report.title.get("revision") or ""
    fills = {"error": "F8D7DA", "warning": "FFF3CD", "info": "DCE8F7"}
    n = 0
    for f in report.findings:
        if f.decision:          # accepted deviations and rejected misreads are not queries
            continue
        n += 1
        size = (page_sizes or {}).get(f.page) if f.page is not None else None
        ws.append([
            f"TQ-{n:03d}", dn, rev, (f.page + 1) if f.page is not None else "",
            zone(f, *size) if size else "", f.rule, f.severity, f.query,
            f.standard + (f" cl. {f.clause}" if f.clause else ""), "", "Open",
        ])
        ws.cell(ws.max_row, 7).fill = PatternFill("solid", fgColor=fills[f.severity])
    for c in ws[1]:
        c.font = Font(bold=True)
    widths = [9, 18, 6, 6, 6, 8, 9, 80, 26, 40, 9]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wdt
    for row in ws.iter_rows(min_row=2):
        row[7].alignment = Alignment(wrap_text=True, vertical="top")

    fs = wb.create_sheet("Findings")
    keys = ["id", "rule", "severity", "title", "page", "evidence", "confidence", "decision", "standard", "message"]
    fs.append(keys)
    for f in report.findings + report.suppressed:
        d = f.to_dict()
        fs.append([d[k] if d[k] is not None else "" for k in keys])
    for c in fs[1]:
        c.font = Font(bold=True)

    ss = wb.create_sheet("Summary")
    for k, v in [("Source", Path(report.source).name), ("Pages", report.pages),
                 ("Errors", report.count("error")), ("Warnings", report.count("warning")),
                 ("Info", report.count("info")), ("Suppressed (accepted deviations)", len(report.suppressed))]:
        ss.append([k, v])
    for k, v in report.stats.items():
        ss.append([k, json.dumps(v) if isinstance(v, (dict, list)) else v])
    wb.save(str(out))
    return out
